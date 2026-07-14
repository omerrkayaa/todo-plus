from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "screenshots"
OUT = ROOT / "ToDo_Plus_Gelistirme_Sureci_Raporu.docx"

NAVY = "17324D"
BLUE = "2563EB"
INK = "17202A"
GRAY = "667085"
LIGHT = "EAF1FB"
PALE = "F4F6F9"
GREEN = "18794E"
AMBER = "9A6700"
WHITE = "FFFFFF"


def set_font(run, size=11, bold=False, color=INK, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, NAVY, 8, 4),
):
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for list_name in ("List Bullet", "List Number"):
    style = doc.styles[list_name]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.375)
    style.paragraph_format.first_line_indent = Inches(-0.194)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.208

header = section.header.paragraphs[0]
set_font(header.add_run("ToDo+ | Geliştirme Süreci Ara Raporu"), 9, True, GRAY)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Ömer Kaya • 2026"), 9, False, GRAY)


def paragraph(text="", *, bold=False, color=INK, italic=False, align=None, after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), 11, bold, color, italic)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def callout(label, text, fill=LIGHT, label_color=BLUE, cant_split=True):
    table = doc.add_table(rows=1, cols=1)
    if cant_split:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(label + "  "), 10.5, True, label_color)
    set_font(p.add_run(text), 10.5)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def status_table(rows):
    table = doc.add_table(rows=1, cols=3)
    headers = ("Çalışma alanı", "Mevcut durum", "Açıklama")
    for index, value in enumerate(headers):
        shade(table.cell(0, index), NAVY)
        set_font(table.cell(0, index).paragraphs[0].add_run(value), 10.5, True, WHITE)
    for area, status, explanation in rows:
        cells = table.add_row().cells
        set_font(cells[0].paragraphs[0].add_run(area), 10, True, NAVY)
        status_color = GREEN if status == "Uygulandı" else AMBER
        set_font(cells[1].paragraphs[0].add_run(status), 10, True, status_color)
        set_font(cells[2].paragraphs[0].add_run(explanation), 10)
    set_table_geometry(table, [2350, 1550, 5460])
    return table


def find_new_shot(stem):
    candidates = sorted(SHOTS.glob(f"{stem}*yeni*.png"))
    if candidates:
        return candidates[0]
    candidates = sorted(SHOTS.glob(f"{stem}-*.png"))
    return candidates[0] if candidates else None


def add_picture(stem, caption):
    path = find_new_shot(stem)
    if path is None:
        marker = paragraph(
            f"[Güncel ekran görüntüsü bekleniyor: {stem}]",
            bold=True,
            color=AMBER,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        marker.paragraph_format.space_before = Pt(12)
        marker.paragraph_format.space_after = Pt(12)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    if stem == "08":
        p.add_run().add_picture(str(path), height=Inches(6.7))
    else:
        p.add_run().add_picture(str(path), width=Inches(6.35))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(3)
    cp.paragraph_format.space_after = Pt(10)
    set_font(cp.add_run(caption), 9.5, False, GRAY, True)


# Kapak
doc.add_paragraph().paragraph_format.space_after = Pt(72)
paragraph("GELİŞTİRME SÜRECİ ARA RAPORU", bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
p = paragraph("ToDo+", bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
for run in p.runs:
    run.font.size = Pt(34)
p = paragraph("Görev Takip Sistemi", bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
for run in p.runs:
    run.font.size = Pt(22)
paragraph(
    "Şu ana kadar geliştirilen modüller, mevcut çalışma durumu ve sonraki adımlar",
    color=GRAY,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    after=72,
)
meta = doc.add_table(rows=4, cols=2)
for row, (label, value) in enumerate((
    ("Hazırlayan", "Ömer Kaya"),
    ("Çalışma", "ToDo+ Görev Takip Sistemi"),
    ("Rapor türü", "Geliştirme süreci / mevcut durum"),
    ("Tarih", date.today().strftime("%d.%m.%Y")),
)):
    shade(meta.cell(row, 0), LIGHT)
    set_font(meta.cell(row, 0).paragraphs[0].add_run(label), 11, True, NAVY)
    set_font(meta.cell(row, 1).paragraphs[0].add_run(value), 11)
set_table_geometry(meta, [2700, 6660])
doc.add_paragraph().paragraph_format.space_after = Pt(26)
paragraph(
    "HTML • CSS • JavaScript • LocalStorage • Supabase/PostgreSQL • Playwright",
    color=GRAY,
    align=WD_ALIGN_PARAGRAPH.CENTER,
)
doc.add_page_break()

doc.add_heading("Raporun amacı", level=1)
doc.add_paragraph(
    "Bu rapor, ToDo+ projesinin tamamlanmış bir ürün olduğunu ileri sürmek için değil; geliştirme sürecinde şu ana kadar ortaya çıkarılan tasarım, özellik, veri modeli ve test çalışmalarını düzenli biçimde göstermek için hazırlanmıştır. Çalışan demo akışları ile henüz gerçek ortam doğrulaması bekleyen bölümler birbirinden ayrılmıştır."
)
callout(
    "Mevcut aşama",
    "Uygulamanın arayüzü ve temel görev yönetimi akışları demo modunda çalışmaktadır. Supabase altyapısı hazırlanmış olmakla birlikte gerçek proje bilgileriyle uçtan uca test çalışmaları devam eden aşama olarak ele alınmaktadır.",
)

doc.add_heading("1. Projenin amacı ve kapsamı", level=1)
doc.add_paragraph(
    "ToDo+, bireysel ve ekip çalışmalarındaki görevlerin proje, kategori, kullanıcı, öncelik, durum ve teslim tarihi bilgileriyle takip edilebilmesini amaçlayan web tabanlı bir görev yönetimi çalışmasıdır. Proje, basit bir yapılacaklar listesinin ötesine geçerek yönetici ve kullanıcı rollerine göre farklı çalışma alanları sunacak biçimde geliştirilmiştir."
)
for text in (
    "Yönetici ve kullanıcı için rol bazlı ekranlar oluşturmak.",
    "Proje, kategori ve görev kayıtlarını yönetebilmek.",
    "Görevleri kişilere atamak ve görevi veren kullanıcıları ilişkilendirmek.",
    "Liste ve Kanban görünümleri üzerinden görev durumunu izlemek.",
    "Yorum, durum geçmişi, teslim tarihi ve başarı oranı gibi destekleyici bilgileri göstermek.",
    "Uygulamayı masaüstü ve mobil ekranlarda kullanılabilir hâle getirmek.",
):
    bullet(text)

doc.add_heading("2. Kullanılan teknolojiler", level=1)
status_table((
    ("Arayüz", "Uygulandı", "HTML5 ve CSS3 ile responsive tek sayfalık kullanıcı arayüzü."),
    ("İstemci mantığı", "Uygulandı", "Vanilla JavaScript ile state, CRUD, filtreleme ve render işlemleri."),
    ("Yerel veri", "Uygulandı", "Demo kullanımında LocalStorage tabanlı kalıcı veri saklama."),
    ("Bulut altyapısı", "Geliştiriliyor", "Supabase Authentication, PostgreSQL şeması ve RLS politikaları."),
    ("Test", "Geliştiriliyor", "Playwright tabanlı temel kullanıcı akışları için smoke test."),
    ("Modelleme", "Uygulandı", "Draw.io ve Mermaid ile ilişkisel veri modeli gösterimi."),
))

doc.add_heading("3. Şu ana kadar geliştirilen çalışmalar", level=1)
doc.add_heading("3.1 Giriş ve kullanıcı rolleri", level=2)
doc.add_paragraph(
    "Giriş ve kayıt ekranları hazırlanmış, demo hesaplarıyla yönetici veya kullanıcı paneline erişim sağlanmıştır. Yönetici; ekip, proje, kategori ve görev yönetimi alanlarına erişebilirken normal kullanıcı yalnızca kendisine atanan görevleri ve kişisel istatistiklerini görmektedir."
)
add_picture("01", "Şekil 1. Güncel giriş ve kayıt ekranı")

doc.add_heading("3.2 Yönetici genel bakış ekranı", level=2)
doc.add_paragraph(
    "Yönetici panelinde toplam, tamamlanan, devam eden ve geciken görev sayıları gösterilmektedir. Kullanıcı başarıları öncelik puanlarına göre hesaplanmakta; yaklaşan teslimler, son görev hareketleri ve ekip iş yükü önerileri aynı ekranda özetlenmektedir."
)
add_picture("02", "Şekil 2. Güncel yönetici genel bakış ekranı")

doc.add_heading("3.3 Proje ve kategori yönetimi", level=2)
doc.add_paragraph(
    "Projeler için oluşturma, listeleme, düzenleme ve silme akışları hazırlanmıştır. Proje kartları görev sayısını ve tamamlanma oranını gösterir. Kategoriler renk bilgisiyle tanımlanabilir ve görevlere atanabilir."
)
add_picture("03", "Şekil 3. Güncel proje ve kategori yönetimi ekranı")

doc.add_heading("3.4 Görev yönetimi", level=2)
doc.add_paragraph(
    "Görev formunda başlık, açıklama, proje, kategori, atanan kullanıcı, görevi veren kullanıcılar, tahmini süre, öncelik, durum ve son teslim tarihi alanları bulunmaktadır. Görevler düzenlenebilir, silinebilir, tamamlanabilir veya yeniden açılabilir."
)
add_picture("04", "Şekil 4. Güncel görev listesi ve filtreler")

doc.add_heading("3.5 Görev detayı, yorum ve geçmiş", level=2)
doc.add_paragraph(
    "Görev detay penceresinde görevin tüm alanları, durum geçmişi ve yorumları birlikte gösterilmektedir. Kullanıcılar detay ekranından yorum ekleyebilir ve yetkileri kapsamında görev durumunu değiştirebilir."
)
add_picture("05", "Şekil 5. Güncel görev detayı ve yorum alanı")

doc.add_heading("3.6 Kanban görünümü", level=2)
doc.add_paragraph(
    "Görevler bekliyor, devam ediyor ve tamamlandı sütunlarında Kanban panosu üzerinde gösterilmektedir. Yönetici sürükle-bırak veya durum kontrolleri aracılığıyla görev akışını güncelleyebilir. Liste ve pano aynı uygulama durumunu kullanmaktadır."
)
add_picture("06", "Şekil 6. Güncel Kanban panosu")

doc.add_heading("3.7 Responsive ve mobil kullanım", level=2)
doc.add_paragraph(
    "Dar ekranlarda kenar menüsü açılır yapıya dönüşmekte; istatistik kartları, görevler ve formlar tek sütunlu düzende gösterilmektedir. Mobil menü arka plan tıklamasıyla kapanabilir ve temel işlemler küçük ekranlarda erişilebilir durumdadır."
)
add_picture("08", "Şekil 7. Güncel mobil genel bakış ekranı")

doc.add_heading("4. Veri modeli ve Supabase hazırlığı", level=1)
doc.add_paragraph(
    "Projede users, projects, categories, tasks, task_assignments, task_givers, task_history ve task_comments tablolarını içeren PostgreSQL şeması hazırlanmıştır. Foreign-key ilişkileri, indeksler, kontrol kısıtları, başarı oranı görünümü ve Row Level Security politikaları SQL dosyasında tanımlanmıştır."
)
callout(
    "Sınır",
    "Bu bölüm altyapı hazırlığını göstermektedir. Gerçek Supabase URL ve anahtarıyla kayıt, giriş, senkronizasyon ve RLS senaryolarının uçtan uca doğrulanması sonraki geliştirme adımıdır.",
    fill="FFF4E5",
    label_color=AMBER,
)

doc.add_heading("5. Başarı oranı ve karar desteği", level=1)
doc.add_paragraph(
    "Başarı oranı yalnızca tamamlanan görev adedine göre değil, görev önceliklerine verilen puanlara göre hesaplanmaktadır. Düşük öncelikli görev 1, orta öncelikli görev 2 ve yüksek öncelikli görev 3 puan değerindedir. Tamamlanan görevlerin puanı toplam atanmış görev puanına bölünerek yüzde değer elde edilir."
)
callout("Formül", "Başarı oranı = tamamlanan görev puanı / atanmış toplam görev puanı × 100")
doc.add_paragraph(
    "Yönetici panelindeki karar desteği bölümü; aktif görev sayısı, geciken görev, tahmini iş yükü ve başarı oranı gibi ölçümleri birlikte değerlendirerek görev ataması için aday kullanıcı önerisi sunmaktadır. Bu özellik kararın yerine geçmez; yöneticiye özet bilgi sağlar."
)

doc.add_heading("6. Test ve doğrulama çalışmaları", level=1)
doc.add_paragraph(
    "Playwright tabanlı smoke test dosyası; demo giriş, proje ve kategori oluşturma, görev oluşturma, birleşik filtreler, görev detayı, yorum, durum değiştirme, mobil menü ve rol bazlı görünürlük gibi temel akışları kontrol edecek biçimde hazırlanmıştır. Testlerin yeniden çalıştırılabilmesi için Node.js, proje bağımlılıkları ve yerel web sunucusu gereklidir."
)
for text in (
    "Demo yönetici hesabıyla giriş",
    "Proje, kategori ve görev oluşturma",
    "Arama ile çoklu filtrelerin birlikte çalışması",
    "Görev detayında yorum ve hızlı durum değişikliği",
    "Mobil menünün açılıp kapanması",
    "Normal kullanıcı panelinde yönetim alanlarının gizlenmesi",
    "Tarayıcı çalışma zamanı hatalarının izlenmesi",
):
    bullet(text)

doc.add_heading("7. Mevcut geliştirme durumu", level=1)
status_table((
    ("Arayüz ve responsive tasarım", "Uygulandı", "Masaüstü ve mobil temel ekranlar hazır."),
    ("Demo oturum ve LocalStorage", "Uygulandı", "Kurulum gerektirmeden gösterilebilir çalışma modu hazır."),
    ("Proje/kategori/görev CRUD", "Uygulandı", "Temel oluşturma, düzenleme ve silme akışları mevcut."),
    ("Kanban, filtre ve istatistik", "Uygulandı", "Listeleme ve takip bileşenleri mevcut."),
    ("Supabase gerçek ortam testi", "Geliştiriliyor", "Gerçek proje bilgileriyle entegrasyon doğrulaması bekliyor."),
    ("Çok kullanıcılı üyelik modeli", "Uygulandı", "Proje üyeliği, davet ve proje bazlı rol akışları temel düzeyde hazır."),
    ("Kodun modüllere ayrılması", "Geliştiriliyor", "Yardımcı modüller ayrıldı; ana uygulama dosyası sadeleştirilecek."),
))

doc.add_heading("8. Sonraki geliştirme adımları", level=1)
for text in (
    "Gerçek Supabase projesi üzerinde kayıt, giriş, CRUD, senkronizasyon ve RLS testlerini tamamlamak.",
    "Hazırlanan proje üyeliği, kullanıcı daveti ve rol yönetimi akışlarını gerçek Supabase ortamında doğrulamak.",
    "Ana JavaScript dosyasını auth, state, render ve görev yönetimi modüllerine daha fazla ayırmak.",
    "Senkronizasyon kuyruğunu ağ kesintileri ve eşzamanlı değişiklikler için güçlendirmek.",
    "Dosya ekleri, alt görevler, bildirimler ve raporlama gibi genişletmeleri değerlendirmek.",
):
    bullet(text)

doc.add_heading("9. Genel değerlendirme", level=1)
doc.add_paragraph(
    "ToDo+ projesinde şu ana kadar işlevsel bir demo, rol bazlı çalışma alanları, kapsamlı görev yönetimi bileşenleri, responsive tasarım, ilişkisel veritabanı şeması ve otomatik test altyapısı ortaya çıkarılmıştır. Mevcut sürüm geliştirme sürecindeki ilerlemeyi gösterebilecek düzeydedir. Bundan sonraki çalışma, hazırlanan altyapının gerçek Supabase ortamında doğrulanması, çok kullanıcılı modelin güçlendirilmesi ve kod yapısının sadeleştirilmesi üzerine yoğunlaşacaktır."
)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
set_font(p.add_run("Sonuç: "), 11, True, BLUE)
set_font(
    p.add_run(
        "Rapor, projenin tamamlandığını değil; çalışan bölümlerini, hazırlanmış altyapısını "
        "ve açık biçimde tanımlanmış sonraki adımlarını göstermektedir."
    ),
    11,
)

doc.core_properties.title = "ToDo+ Geliştirme Süreci Ara Raporu"
doc.core_properties.subject = "ToDo+ projesinde şu ana kadar yapılan çalışmalar"
doc.core_properties.author = "Ömer Kaya"
doc.core_properties.keywords = "ToDo+, görev takip, geliştirme süreci, ara rapor"
doc.save(OUT)
print(OUT)
