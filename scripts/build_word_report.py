from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "ToDo_Plus_Proje_Raporu.docx"
SHOT = ROOT / "screenshots"

NAVY = "17324D"
BLUE = "2563EB"
LIGHT = "EAF1FB"
PALE = "F4F6F9"
GRAY = "667085"
LINE = "D7DEE8"
WHITE = "FFFFFF"
BLACK = "17202A"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(1)
sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

def font(run, size=11, bold=False, color=BLACK, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, NAVY, 8, 4),
]:
    s = styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

for list_name in ["List Bullet", "List Number"]:
    s = styles[list_name]
    s.font.name = "Calibri"
    s.font.size = Pt(11)
    s.paragraph_format.left_indent = Inches(0.375)
    s.paragraph_format.first_line_indent = Inches(-0.194)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.line_spacing = 1.208

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = OxmlElement("w:tblInd")
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    tblPr.append(tblInd)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(width))
        grid.append(gc)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tcW.set(qn("w:w"), str(widths[idx]))
            tcW.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def add_header_footer(section):
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font(hp.add_run("ToDo+ Görev Takip Sistemi | Proje Raporu"), 9, True, GRAY)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(fp.add_run("Ömer Kaya • 2026"), 9, False, GRAY)

add_header_footer(sec)

def title(text, size, color=NAVY, after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    font(p.add_run(text), size, True, color)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p

custom_numbering_counter = 10

def new_numbering():
    global custom_numbering_counter
    numbering = doc.part.numbering_part.element
    num_id = custom_numbering_counter
    custom_numbering_counter += 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    aid = OxmlElement("w:abstractNumId"); aid.set(qn("w:val"), "7")
    override = OxmlElement("w:lvlOverride"); override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride"); start.set(qn("w:val"), "1")
    override.append(start)
    num.extend([aid, override])
    numbering.append(num)
    return num_id

def number(text, num_id):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id))
    numPr.extend([ilvl, nid]); pPr.append(numPr)
    p.add_run(text)
    return p

def callout(label, text, fill=LIGHT):
    t = doc.add_table(rows=1, cols=1)
    set_table_geometry(t, [9360])
    cell = t.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run(label + "  "), 10.5, True, BLUE)
    font(p.add_run(text), 10.5, False, BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_picture(filename, caption):
    path = SHOT / filename
    if not path.exists(): return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.35))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(3)
    cp.paragraph_format.space_after = Pt(10)
    font(cp.add_run(caption), 9.5, False, GRAY, True)

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(72)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("PROJE RAPORU"), 12, True, BLUE)
title("ToDo+", 34, NAVY, 4)
title("Görev Takip Sistemi", 22, BLUE, 18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Proje, görev, kullanıcı ve başarı yönetimi için web tabanlı uygulama"), 13, False, GRAY)
doc.add_paragraph().paragraph_format.space_after = Pt(95)
meta = doc.add_table(rows=3, cols=2)
set_table_geometry(meta, [2700, 6660])
for i, (k, v) in enumerate([("Hazırlayan", "Ömer Kaya"), ("Çalışma", "ToDo+ Görev Takip Sistemi"), ("Yıl", "2026")]):
    shade(meta.cell(i, 0), LIGHT)
    font(meta.cell(i, 0).paragraphs[0].add_run(k), 11, True, NAVY)
    font(meta.cell(i, 1).paragraphs[0].add_run(v), 11, False, BLACK)
doc.add_paragraph().paragraph_format.space_after = Pt(28)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("HTML • CSS • JavaScript • Supabase/PostgreSQL • Playwright"), 10, False, GRAY)
doc.add_page_break()

doc.add_heading("İçindekiler", level=1)
for item in [
    "1. Yönetici Özeti", "2. Projenin Amacı ve Kapsamı", "3. Kullanılan Teknolojiler",
    "4. Uygulama Mimarisi", "5. Özellikler ve Kullanıcı Akışları", "6. Veritabanı Tasarımı",
    "7. Başarı Oranı ve İstatistikler", "8. Ekranlar ve Kullanım", "9. Supabase Entegrasyonu",
    "10. Test ve Kalite Güvencesi", "11. Güvenlik ve Bilinen Sınırlar", "12. Sonuç ve Yol Haritası"
]: bullet(item)

doc.add_heading("1. Yönetici Özeti", level=1)
doc.add_paragraph("ToDo+, klasik yapılacaklar listelerinden daha kapsamlı bir görev takip sistemi olarak geliştirilmiştir. Kullanıcılar proje ve kategori oluşturabilir, görevleri kişilere atayabilir, öncelik ve teslim tarihi belirleyebilir, Kanban panosunda durum değiştirebilir ve görev detayında yorum yazabilir. Uygulama ayrıca görev geçmişini tutar ve kullanıcı başarısını görev önceliklerine göre ağırlıklı biçimde hesaplar.")
callout("Öne çıkan değer", "Uygulama kurulum gerektirmeyen LocalStorage tabanlı demo moduyla çalışır; yapılandırıldığında Supabase Authentication ve PostgreSQL tablolarına bağlanabilir.")

doc.add_heading("2. Projenin Amacı ve Kapsamı", level=1)
doc.add_paragraph("Projenin amacı, bireysel veya ekip çalışmalarındaki görevlerin planlanması, sorumlulara atanması, durumlarının izlenmesi ve sonuçların ölçülebilmesi için anlaşılır bir web arayüzü sunmaktır. Tasarım; eğitim projesi, portföy çalışması ve daha gelişmiş ekip yönetimi ürünleri için başlangıç altyapısı olarak kullanılabilecek biçimde hazırlanmıştır.")
for text in ["Proje, kategori ve görevler için CRUD işlemleri", "Kullanıcı, görevi veren ve atanan kişi ilişkileri", "Arama, çoklu filtreleme ve sıralama", "Kanban ve görev listesi görünümleri", "Yorum, geçmiş ve teslim tarihi takibi", "Öncelik katsayılı başarı raporlaması", "Mobil uyum ve erişilebilir klavye kullanımı"]: bullet(text)

doc.add_heading("3. Kullanılan Teknolojiler", level=1)
table = doc.add_table(rows=1, cols=3)
set_table_geometry(table, [2200, 2860, 4300])
repeat_header(table.rows[0])
for i, h in enumerate(["Katman", "Teknoloji", "Kullanım amacı"]):
    shade(table.cell(0, i), NAVY); font(table.cell(0, i).paragraphs[0].add_run(h), 10.5, True, WHITE)
rows = [
    ("Arayüz", "HTML5 / CSS3", "Semantik sayfa yapısı, responsive düzen ve görsel bileşenler"),
    ("İstemci", "JavaScript", "State, CRUD, filtreleme, render ve kullanıcı etkileşimleri"),
    ("Yerel veri", "LocalStorage", "Demo modunda kalıcı tarayıcı verisi"),
    ("Bulut", "Supabase", "Authentication, PostgreSQL ve RLS"),
    ("Test", "Playwright", "Temel kullanıcı akışlarının tarayıcı otomasyonu"),
    ("Modelleme", "Draw.io / Mermaid", "ER diyagramı ve veri ilişkileri"),
]
for r in rows:
    cells = table.add_row().cells
    for i, value in enumerate(r): font(cells[i].paragraphs[0].add_run(value), 10)
set_table_geometry(table, [2200, 2860, 4300])

doc.add_heading("4. Uygulama Mimarisi", level=1)
doc.add_paragraph("Uygulama istemci tarafında çalışan tek sayfalı bir yapıya sahiptir. index.html ekran iskeletini, styles.css tasarım sistemini, app.js ise uygulama davranışını yönetir. State bellekte tutulur ve çalışma moduna göre LocalStorage veya Supabase ile eşitlenir.")
for text in ["Authentication: demo hesapları veya Supabase Auth oturumu", "State: kullanıcılar, projeler, kategoriler, görevler, atamalar, verenler, geçmiş ve yorumlar", "Render: state değiştikçe istatistik, proje, görev ve Kanban görünümünün yenilenmesi", "Persistence: yerel kayıt ve yapılandırılmışsa kuyruklanmış Supabase senkronizasyonu", "Security: ilişkili görevlere göre Row Level Security politikaları"]: bullet(text)

doc.add_heading("5. Özellikler ve Kullanıcı Akışları", level=1)
doc.add_heading("5.1 Kullanıcı ve oturum yönetimi", level=2)
doc.add_paragraph("Kullanıcı giriş veya kayıt ekranından uygulamaya erişir. Demo modunda varsayılan hesap kullanılabilir; Supabase modu etkinse gerçek e-posta ve parola doğrulaması yapılır. Arayüz bağlantı durumunu kullanıcıya gösterir.")
doc.add_heading("5.2 Proje ve kategori yönetimi", level=2)
doc.add_paragraph("Projeler görevlerin ana çalışma alanıdır. Her proje için açıklama, oluşturulma tarihi, görev sayısı ve tamamlanma yüzdesi gösterilir. Kategoriler renk bilgisiyle görevleri sınıflandırır.")
doc.add_heading("5.3 Görev yönetimi", level=2)
doc.add_paragraph("Görev formunda başlık, açıklama, proje, kategori, atanan kişi, görevi veren kişiler, tahmini süre, öncelik, durum ve son teslim tarihi bulunur. Görevler liste veya Kanban üzerinde görüntülenebilir. Detay modalında tüm bilgiler, geçmiş ve yorumlar birlikte sunulur.")
doc.add_heading("5.4 Filtreleme ve kullanılabilirlik", level=2)
doc.add_paragraph("Görevler metin, durum, öncelik, kategori ve atanan kullanıcıya göre filtrelenebilir. Sonuçlar teslim tarihine veya önceliğe göre sıralanabilir. Filtre temizleme işlemi tüm seçimleri başlangıç değerlerine döndürür.")

doc.add_heading("6. Veritabanı Tasarımı", level=1)
doc.add_paragraph("Supabase şeması ilişkisel PostgreSQL modeli kullanır. UUID anahtarları, foreign-key kuralları, kontrol kısıtları, indeksler, başarı oranı görünümü ve RLS politikaları supabase-schema.sql içinde tanımlanmıştır.")
table = doc.add_table(rows=1, cols=2)
set_table_geometry(table, [2600, 6760])
repeat_header(table.rows[0])
for i, h in enumerate(["Tablo", "Görevi"]): shade(table.cell(0,i), NAVY); font(table.cell(0,i).paragraphs[0].add_run(h), 10.5, True, WHITE)
for name, purpose in [("users", "Kullanıcı profil bilgileri"), ("projects", "Kullanıcı projeleri"), ("categories", "Görev sınıflandırmaları"), ("tasks", "Ana görev kayıtları"), ("task_assignments", "Atanan kullanıcı ve kazanılan puan"), ("task_givers", "Görevi veren kişiler"), ("task_history", "Durum değişikliği kayıtları"), ("task_comments", "Görev yorumları")]:
    cells = table.add_row().cells; font(cells[0].paragraphs[0].add_run(name), 10, True, BLUE); font(cells[1].paragraphs[0].add_run(purpose), 10)
set_table_geometry(table, [2600, 6760])
doc.add_heading("6.1 Temel ilişkiler", level=2)
for text in ["Bir kullanıcı birden fazla proje ve kategori oluşturabilir.", "Bir proje birden fazla görev içerir.", "Bir görev bir kullanıcıya atanır ve birden fazla kişi tarafından verilebilir.", "Bir görevin birden fazla geçmiş ve yorum kaydı olabilir.", "Proje veya görev silindiğinde bağlı kayıtlar veritabanı kurallarına göre temizlenir."]: bullet(text)

doc.add_heading("7. Başarı Oranı ve İstatistikler", level=1)
doc.add_paragraph("Başarı oranı yalnızca tamamlanan görev adedine göre değil, görev önceliğine göre hesaplanır. Böylece yüksek öncelikli bir görevin tamamlanması sonuca daha fazla katkıda bulunur.")
callout("Formül", "Başarı oranı = tamamlanan görev puanları / atanmış tüm görev puanları × 100")
for text in ["Düşük öncelik: 1 puan", "Orta öncelik: 2 puan", "Yüksek öncelik: 3 puan"]: bullet(text)
doc.add_paragraph("Örneğin yüksek, orta ve düşük öncelikli üç görev toplam 6 puandır. Kullanıcı yüksek ve düşük öncelikli görevleri tamamladığında 4 puan kazanır ve başarı oranı yaklaşık %66,7 olur.")

doc.add_heading("8. Ekranlar ve Kullanım", level=1)
doc.add_heading("8.1 Giriş ve kayıt ekranı", level=2)
doc.add_paragraph("Uygulama ilk açıldığında kullanıcıyı sade bir giriş ekranı karşılar. Giriş ve kayıt sekmeleri aynı kart üzerinde sunulur. Demo hesabı bilgileri ekranın altında belirtilir.")
add_picture("01-giris-ekrani.png", "Şekil 1. ToDo+ giriş ekranı")
doc.add_heading("8.2 İstatistik ekranı", level=2)
doc.add_paragraph("İstatistik görünümü toplam, tamamlanan, devam eden ve geciken görev sayılarını; başarı oranını, yaklaşan teslimleri ve son görev hareketlerini gösterir.")
add_picture("02-istatistik-ekrani.png", "Şekil 2. İstatistik ve başarı görünümü")
doc.add_heading("8.3 Proje ve kategori ekranı", level=2)
doc.add_paragraph("Projeler ve kategoriler aynı çalışma alanında yönetilir. Proje kartlarında açıklama, görev sayısı, tamamlanan görev sayısı ve ilerleme yüzdesi bulunur.")
add_picture("03-projeler-kategoriler.png", "Şekil 3. Proje ve kategori yönetimi")
doc.add_heading("8.4 Görev ve Kanban ekranı", level=2)
doc.add_paragraph("Görev ekranı filtreleri, Kanban panosunu, görev formunu ve görev listesini bir araya getirir. Aynı veri farklı görünümlerde eş zamanlı olarak güncellenir.")
add_picture("04-gorevler-ekrani.png", "Şekil 4. Görev listesi, filtreler ve Kanban panosu")
doc.add_heading("8.5 Görev detayı ve yorumlar", level=2)
doc.add_paragraph("Detay modalı görev bilgilerini, kullanıcı ilişkilerini, geçmiş kayıtlarını ve yorumları tek alanda gösterir. Kullanıcı modal üzerinden yeni yorum ekleyebilir.")
add_picture("05-gorev-detayi-yorumlar.png", "Şekil 5. Görev detayı, geçmiş ve yorum alanı")

doc.add_heading("9. Supabase Entegrasyonu", level=1)
doc.add_paragraph("Supabase bilgileri tanımlandığında uygulama gerçek Authentication oturumunu kullanır. Kullanıcı profili users tablosuna yazılır; projeler, kategoriler, görevler ve bağlı kayıtlar UUID değerleriyle senkronize edilir. Senkronizasyon durumu arayüzde yükleniyor, kaydediliyor, güncel veya hata olarak gösterilir.")
doc.add_heading("9.1 Kurulum özeti", level=2)
setup_num = new_numbering()
for text in ["Supabase projesi oluşturulur.", "supabase-schema.sql SQL Editor içinde çalıştırılır.", "Project URL ve anon key supabase-config.js içine yazılır.", "Email Authentication etkinleştirilir.", "Kayıt, giriş, CRUD ve RLS senaryoları gerçek hesaplarla test edilir."]: number(text, setup_num)
callout("Güvenlik", "Tarayıcıya yalnızca anon/public key yazılmalıdır. service_role anahtarı istemci kodunda kesinlikle kullanılmamalıdır.", "FFF4E5")

doc.add_heading("10. Test ve Kalite Güvencesi", level=1)
doc.add_paragraph("Proje tesliminden önce JavaScript sözdizimi, Git diff temizliği ve tarayıcı tabanlı smoke test uygulanmıştır. Otomatik test sekiz temel kullanıcı akışını doğrular.")
for text in ["Demo hesabıyla giriş", "Proje ve kategori oluşturma", "Görev oluşturma", "Arama ve filtre temizleme", "Görev detayını açma ve yorum ekleme", "Görev durumunu tamamlama", "Mobil menüyü açma", "Tarayıcı çalışma zamanı hatalarını kontrol etme"]: bullet(text)
doc.add_heading("10.1 Manuel kabul kriterleri", level=2)
for text in ["CRUD işlemleri listeleri ve istatistikleri anında güncellemelidir.", "Kanban ve görev listesi aynı state değerini göstermelidir.", "Geciken görev hesabı tamamlanmış görevleri dışarıda bırakmalıdır.", "Escape tuşu modalı kapatmalı, mobil menü dar ekranda çalışmalıdır.", "Supabase ağ hatasında yerel veri korunmalı ve kullanıcı bilgilendirilmelidir."]: bullet(text)

doc.add_heading("11. Güvenlik ve Bilinen Sınırlar", level=1)
doc.add_paragraph("SQL şeması RLS politikaları içerir; görev oluşturucusu, atanan kullanıcı ve görevi veren kullanıcı erişimi merkezi bir fonksiyonla değerlendirilir. Bununla birlikte canlı kullanımdan önce gerçek Supabase ortamında entegrasyon ve yetki testleri tamamlanmalıdır.")
for text in ["Demo kullanıcı ekleme gerçek Supabase Auth hesabı oluşturmaz.", "Demo verilerini sıfırlama yalnızca demo modunda kullanılmalıdır.", "Kullanıcı profil ve e-posta görünürlüğü üretim gereksinimine göre daraltılmalıdır.", "Test scriptlerindeki makineye özel yollar taşınabilir hâle getirilmelidir.", "Senkronizasyon kuyruğu yoğun eş zamanlı değişiklikler için ayrıca güçlendirilmelidir.", "Proje üyeliği ve rol modeli henüz ayrı tablolarla uygulanmamıştır."]: bullet(text)

doc.add_heading("12. Sonuç ve Yol Haritası", level=1)
doc.add_paragraph("ToDo+ projesi, görev yönetiminin temel ve orta seviye gereksinimlerini çalışan bir web arayüzünde birleştirmiştir. Demo modunda hızlıca değerlendirilebilir; Supabase şeması ve entegrasyon katmanı sayesinde gerçek kullanıcı ve veritabanı altyapısına geçiş için temel sunar. Kanban, yorum, geçmiş, filtreleme ve ağırlıklı başarı puanı projeyi basit bir yapılacaklar listesinin ötesine taşır.")
doc.add_heading("12.1 Önerilen sonraki çalışmalar", level=2)
roadmap_num = new_numbering()
for text in ["Proje üyeliği, davet ve rol tabanlı yetki modeli", "Gerçek Supabase ortamında otomatik RLS testleri", "app.js dosyasının auth, state, render ve veri modüllerine ayrılması", "Alt görevler, dosya ekleri ve Supabase Storage", "Gerçek zamanlı güncelleme ve bildirimler", "Haftalık/aylık grafikler ve gelişmiş raporlama"]: number(text, roadmap_num)
callout("Genel değerlendirme", "Proje eğitim ve portföy teslimi için güçlü bir kapsam sunmaktadır. Canlı üretim kullanımı öncesinde güvenlik, kullanıcı yönetimi ve eş zamanlı senkronizasyon maddeleri tamamlanmalıdır.")

doc.core_properties.title = "ToDo+ Görev Takip Sistemi Proje Raporu"
doc.core_properties.subject = "Web tabanlı görev takip sistemi teknik ve kullanıcı raporu"
doc.core_properties.author = "Ömer Kaya"
doc.core_properties.keywords = "ToDo+, görev takip, Supabase, JavaScript, proje raporu"
doc.save(OUT)
print(OUT)
